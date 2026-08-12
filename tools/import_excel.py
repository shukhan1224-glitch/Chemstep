#!/usr/bin/env python3
"""
把旧的 Excel 点名表(矩阵式:横轴日期、纵轴学生、1=出席 0=缺席)
转成新系统可以直接贴上的 CSV。

用法:
    python3 tools/import_excel.py 旧点名表.xlsx -o out/
    python3 tools/import_excel.py 旧点名表.xlsx -o out/ \
        --class-fee 初二科学=80/4 --class-fee 高三=70/4
    python3 tools/import_excel.py 旧点名表.xlsx -o out/ --history none

产生:
    out/学生.csv        贴进「学生」表
    out/课程记录.csv     贴进「课程记录」表   (--history all 时才有)
    out/付款记录.csv     贴进「付款记录」表
    out/汇入报告.txt     做了什么、跳过了什么、每个人现在欠几堂

── 两个会自动处理的旧档特性 ──────────────────────────────────────
1. 日期日月对调
   旧档的日期栏常常一半是文字、一半是数字。原因是 Excel 的地区设定是
   「月/日/年」而使用者输入「日/月/年」:
       输入 19/1/2026 → 19 不可能是月份 → Excel 放弃解析 → 存成文字(正确)
       输入 12/1/2026 → 当成 12月1日   → 存成数字(日月被对调了)
   预设会把数字那一半的日月对调回来,并用「同一班的星期是否一致」验证。
   要关掉请加 --no-swap-fix。

2. 底色 = 已缴费
   很多人会把已经收过钱的堂数涂上颜色(通常是黄色),从第一堂往后连续涂。
   预设会自动找出档案里最常出现的那个底色,把涂色的格子数当成「已付堂数」,
   转成付款记录。要指定颜色用 --paid-fill FFFF00,要关掉用 --paid-fill none。
"""

import argparse
import csv
import datetime
import os
import sys
from collections import Counter, OrderedDict

try:
    import openpyxl
except ImportError:
    sys.exit('需要 openpyxl:pip install openpyxl')

EPOCH = datetime.date(1899, 12, 30)
WEEK = '一二三四五六日'

# 旧档的 1 / 0 是「这堂算不算钱」,不是「有没有出现」:
#   1 = 算钱     → 出席
#   0 = 不算钱   → 请假(免扣)(老师通融的那几堂)
# 少数人的旧档 0 是「缺席但照算」,那种用 --zero-charged 切换。
STATUS_PRESENT = '出席'
STATUS_ABSENT_FREE = '请假(免扣)'
STATUS_ABSENT_CHARGED = '缺席(照算)'
YES, NO = '是', '否'
CHARGEABLE = {STATUS_PRESENT: YES, STATUS_ABSENT_CHARGED: YES, STATUS_ABSENT_FREE: NO}

# 视为「没有底色」的值
BLANK_FILLS = {'None', '00000000', 'FFFFFFFF', 'FFFFFF'}

STUDENT_HEADERS = ['学生ID', '姓名', '班级', '家长称呼', 'WhatsApp号码',
                   '每堂收费RM', '配套堂数', '状态', '备注']
SESSION_HEADERS = ['记录ID', '日期', '班级', '学生ID', '学生姓名',
                   '出席状态', '是否扣堂', '备注', '记录时间', '最后修改']
PAYMENT_HEADERS = ['付款ID', '日期', '学生ID', '学生姓名', '金额RM',
                   '购买堂数', '付款方式', '备注', '记录时间']


# ── 日期解析 ────────────────────────────────────────────────────────

def from_serial(n):
    return EPOCH + datetime.timedelta(days=int(n))


def swap_day_month(d):
    """2026-12-01 → 2026-01-12。日或月超过 12 就换不了,回传 None。"""
    try:
        return datetime.date(d.year, d.day, d.month)
    except ValueError:
        return None


def parse_header_date(value, swap_fix):
    """回传 (date, 来源说明) 或 (None, 原因)"""
    if value is None:
        return None, '空白'

    if isinstance(value, datetime.datetime):
        return value.date(), '真日期'
    if isinstance(value, datetime.date):
        return value, '真日期'

    if isinstance(value, (int, float)):
        raw = from_serial(value)
        if swap_fix:
            fixed = swap_day_month(raw)
            if fixed and fixed != raw:
                return fixed, '数字%d,日月对调 %s → %s' % (int(value), raw, fixed)
            return raw, '数字%d(日月相同或无法对调)' % int(value)
        return raw, '数字%d,未修正' % int(value)

    text = str(value).strip()
    for fmt in ('%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d', '%d-%m-%Y'):
        try:
            return datetime.datetime.strptime(text, fmt).date(), '文字「%s」' % text
        except ValueError:
            continue
    return None, '看不懂的日期「%s」' % text


# ── 底色 ────────────────────────────────────────────────────────────

def fill_rgb(cell):
    """回传储存格实心底色的 RGB 字串;没有底色回传 None"""
    f = cell.fill
    if f is None or f.fill_type != 'solid':
        return None
    rgb = str(getattr(f.fgColor, 'rgb', None))
    if rgb in BLANK_FILLS:
        return None
    return rgb


def detect_paid_fill(wb_fmt):
    """自动找出「已缴费」的底色:出现次数最多的那个非白色实心底色"""
    counts = Counter()
    for ws in wb_fmt.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                rgb = fill_rgb(cell)
                if rgb:
                    counts[rgb] += 1
    if not counts:
        return None, counts
    return counts.most_common(1)[0][0], counts


# ── 读取旧档 ────────────────────────────────────────────────────────

def find_blocks(ws):
    """
    旧档写满一整列后会另起一个区块继续。
    区块的特徵:A 栏是空的,但右边有日期。
    """
    return [row[0].row for row in ws.iter_rows()
            if row[0].value is None and any(c.value is not None for c in row[1:])]


def read_sheet(ws_val, ws_fmt, swap_fix, paid_fill, zero_charged, report):
    """
    回传 (records, weekdays, bad_dates)
    records = {学生姓名: [(date, 出席状态, 是否已缴), ...]},保留原表的学生顺序
    """
    zero_status = STATUS_ABSENT_CHARGED if zero_charged else STATUS_ABSENT_FREE
    records = OrderedDict()
    weekdays = Counter()
    bad_dates = []

    for block_row in find_blocks(ws_val):
        header = {}
        for cell in ws_val[block_row][1:]:
            if cell.value is None:
                continue
            d, why = parse_header_date(cell.value, swap_fix)
            if d is None:
                bad_dates.append('%s!%s %s' % (ws_val.title, cell.coordinate, why))
            else:
                header[cell.column] = d
                weekdays[WEEK[d.weekday()]] += 1

        r = block_row + 1
        while r <= ws_val.max_row and ws_val.cell(r, 1).value is not None:
            name = str(ws_val.cell(r, 1).value).strip()
            records.setdefault(name, [])
            for col, d in header.items():
                v = ws_val.cell(r, col).value
                paid = paid_fill is not None and fill_rgb(ws_fmt.cell(r, col)) == paid_fill
                if v is None or str(v).strip() == '':
                    if paid:
                        report.append('  ⚠️ %s!%s 有底色但没有出席记录,已忽略'
                                      % (ws_val.title, ws_val.cell(r, col).coordinate))
                    continue          # 空白 = 当时还没入学 / 已退出,不是缺席
                try:
                    n = int(float(v))
                except (TypeError, ValueError):
                    report.append('  ⚠️ %s!%s 看不懂的出席值「%s」,已跳过'
                                  % (ws_val.title, ws_val.cell(r, col).coordinate, v))
                    continue
                status = STATUS_PRESENT if n else zero_status
                # 只有「算钱」的那几堂才算进已缴 —— 老师涂色时会把整段涂满,
                # 包含中间不算钱的 0,那些不该被当成缴过的堂数。
                records[name].append((d, status, paid and CHARGEABLE[status] == YES))
            r += 1

    return records, weekdays, bad_dates


# ── 收费设定 ────────────────────────────────────────────────────────

def parse_class_fee(spec):
    """把 '初二科学=80/4' 解析成 ('初二科学', 80.0, 4)"""
    if '=' not in spec:
        raise argparse.ArgumentTypeError('格式应为 班级=金额/堂数,例如 初二科学=80/4')
    klass, rest = spec.split('=', 1)
    if '/' not in rest:
        raise argparse.ArgumentTypeError('格式应为 班级=金额/堂数,例如 初二科学=80/4')
    amount, lessons = rest.split('/', 1)
    try:
        return klass.strip(), float(amount), int(lessons)
    except ValueError:
        raise argparse.ArgumentTypeError('金额和堂数要是数字,例如 初二科学=80/4')


def fmt_money(x):
    """20.0 → '20';17.5 → '17.50'"""
    return str(int(x)) if float(x).is_integer() else '%.2f' % x


# ── 主流程 ──────────────────────────────────────────────────────────

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx', help='旧的 Excel 点名表')
    ap.add_argument('-o', '--out', default='out', help='输出资料夹(预设 out/)')
    ap.add_argument('--history', choices=['all', 'none'], default='all',
                    help='all=连历史出席记录一起汇入(预设);none=只汇入学生名单')
    ap.add_argument('--no-swap-fix', action='store_true',
                    help='不要修正日月对调(只在你确定旧档日期是对的时候用)')
    ap.add_argument('--paid-fill', default='auto',
                    help='「已缴费」的底色 RGB(例如 FFFF00);auto=自动侦测;none=不读底色')
    ap.add_argument('--zero-charged', action='store_true',
                    help='旧档的 0 是「缺席但照算」而不是「不算钱」时加这个')
    ap.add_argument('--paid-override', action='append', default=[], metavar='姓名=堂数',
                    help='手动指定某位学生的已缴堂数,盖过底色算出来的结果。可重复')
    ap.add_argument('--fee', type=float, default=50, help='每期收费,预设 50')
    ap.add_argument('--package', type=int, default=4, help='每期堂数,预设 4')
    ap.add_argument('--class-fee', action='append', default=[], type=parse_class_fee,
                    metavar='班级=金额/堂数',
                    help='个别班级的收费,可重复,例如 --class-fee 高三=70/4')
    args = ap.parse_args()

    swap_fix = not args.no_swap_fix
    # 两次载入:一次拿值(公式会算好),一次拿格式(底色)
    wb_val = openpyxl.load_workbook(args.xlsx, data_only=True)
    wb_fmt = openpyxl.load_workbook(args.xlsx)
    os.makedirs(args.out, exist_ok=True)

    # 决定「已缴费」的底色
    fill_counts = Counter()
    if args.paid_fill.lower() == 'none':
        paid_fill = None
        fill_note = '关闭(不读底色)'
    elif args.paid_fill.lower() == 'auto':
        paid_fill, fill_counts = detect_paid_fill(wb_fmt)
        fill_note = ('自动侦测到 %s(%d 格)' % (paid_fill, fill_counts[paid_fill])
                     if paid_fill else '档案里没有任何底色,视为没有缴费资料')
    else:
        paid_fill = args.paid_fill.upper()
        if len(paid_fill) == 6:
            paid_fill = 'FF' + paid_fill
        fill_note = '指定为 %s' % paid_fill

    fee_by_class = {k: (amt, n) for k, amt, n in args.class_fee}

    overrides = {}
    for spec in args.paid_override:
        if '=' not in spec:
            sys.exit('--paid-override 的格式是 姓名=堂数,例如 --paid-override 蓝子健=16')
        who, n = spec.split('=', 1)
        try:
            overrides[who.strip()] = int(n)
        except ValueError:
            sys.exit('--paid-override 的堂数要是整数:%s' % spec)
    used_overrides = set()

    report = ['旧 Excel 汇入报告',
              '来源:%s' % os.path.basename(args.xlsx),
              '产生时间:%s' % datetime.datetime.now().strftime('%Y-%m-%d %H:%M'),
              '日月对调修正:%s' % ('开启' if swap_fix else '关闭'),
              '已缴费底色:%s' % fill_note,
              '1 / 0 的解读:1 = 算钱(出席)、0 = %s'
              % ('缺席(照算)' if args.zero_charged else '不算钱 → 请假(免扣)'),
              '历史记录:%s' % ('一起汇入' if args.history == 'all' else '不汇入'),
              '']

    if len(fill_counts) > 1:
        report.append('⚠️ 档案里有不只一种底色,只有最常见的那个被当成「已缴费」:')
        for rgb, n in fill_counts.most_common():
            mark = '  ← 用这个' if rgb == paid_fill else ''
            report.append('     %s  %d 格%s' % (rgb, n, mark))
        report.append('   不对的话用 --paid-fill 指定。')
        report.append('')

    students, sessions, payments = [], [], []
    balances = []
    seen_names = {}
    sid = pid = lid = 0
    today = datetime.date.today()

    for ws_val in wb_val.worksheets:
        klass = ws_val.title.strip()
        ws_fmt = wb_fmt[ws_val.title]
        records, weekdays, bad_dates = read_sheet(ws_val, ws_fmt, swap_fix, paid_fill,
                                                  args.zero_charged, report)
        if not records:
            report.append('【%s】没有资料,跳过' % klass)
            continue

        amount, lessons = fee_by_class.get(klass, (args.fee, args.package))
        per_lesson = amount / lessons if lessons else 0

        main_day, main_count = weekdays.most_common(1)[0]
        total_dates = sum(weekdays.values())
        report.append('【%s】%d 人 · %d 堂课 · 主要在星期%s(%d/%d) · %s堂 RM%s(每堂 RM%s)'
                      % (klass, len(records), total_dates, main_day, main_count, total_dates,
                         lessons, fmt_money(amount), fmt_money(per_lesson)))
        if main_count < total_dates:
            others = ', '.join('星期%s×%d' % (d, n) for d, n in weekdays.most_common()[1:])
            report.append('    其余不在固定日:%s(可能是补课)' % others)
        for b in bad_dates:
            report.append('    ⚠️ %s' % b)

        for name, entries in records.items():
            sid += 1
            student_id = 'S%03d' % sid
            if name in seen_names:
                report.append('    ⚠️ 「%s」在「%s」和「%s」都出现,已当成两个不同的学生'
                              % (name, seen_names[name], klass))
            seen_names[name] = klass

            entries_by_date = sorted(entries, key=lambda x: x[0])
            charged = sum(1 for _, st, _ in entries if CHARGEABLE.get(st) == YES)
            paid_count = sum(1 for _, _, paid in entries if paid)

            if name in overrides:
                report.append('    ✏️ %s 的已缴堂数由 %d 手动改为 %d'
                              % (name, paid_count, overrides[name]))
                paid_count = overrides[name]
                used_overrides.add(name)

            balance = paid_count - charged
            # 老师是「涂满 4 个 1」才算一期,所以已缴一定是配套的整数倍。
            # 不是的话,把上下两个整期的边界算出来让老师挑。
            odd = bool(lessons) and paid_count % lessons != 0
            hint = None
            if odd:
                # entries 依工作表的栏位顺序排列,也就是老师涂色的顺序
                ones = [d for d, st, _ in entries if CHARGEABLE.get(st) == YES]
                low = paid_count - paid_count % lessons
                high = low + lessons
                hint = (low, ones[low - 1] if 0 < low <= len(ones) else None,
                        high, ones[high - 1] if 0 < high <= len(ones) else None)
            balances.append((klass, name, paid_count, charged, balance, odd, lessons, hint))

            students.append([student_id, name, klass, '', '',
                             fmt_money(per_lesson), lessons, '在读',
                             '旧 Excel 汇入:已上 %d 堂、已缴 %d 堂' % (charged, paid_count)])

            if args.history == 'all':
                for d, st, _ in entries_by_date:
                    lid += 1
                    sessions.append(['L%04d' % lid, d.isoformat(), klass,
                                     student_id, name, st, CHARGEABLE.get(st, NO),
                                     '旧 Excel 汇入', '', ''])

            if paid_count:
                pid += 1
                first_paid = min(d for d, _, p in entries if p)
                payments.append(['P%04d' % pid, first_paid.isoformat(), student_id, name,
                                 fmt_money(per_lesson * paid_count), paid_count, '旧档结转',
                                 '旧 Excel 已涂色的 %d 堂,合并成一笔' % paid_count, ''])

    def write_csv(fname, headers, rows):
        with open(os.path.join(args.out, fname), 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f)
            w.writerow(headers)
            w.writerows(rows)

    write_csv('学生.csv', STUDENT_HEADERS, students)
    if sessions:
        write_csv('课程记录.csv', SESSION_HEADERS, sessions)
    if payments:
        write_csv('付款记录.csv', PAYMENT_HEADERS, payments)

    # ── 结算表 ──
    report += ['', '=' * 62,
               '汇入后每个人的余额(余额 = 已缴堂数 − 已上堂数)',
               '=' * 62,
               '%-10s %-12s %5s %5s %6s  %s' % ('班级', '学生', '已缴', '已上', '余额', '状态')]
    owed = 0
    odd_ones = []
    for klass, name, paid_count, charged, balance, odd, lessons, hint in sorted(
            balances, key=lambda x: (x[0], x[4])):
        tag = ('🔴 欠 %d 堂' % -balance if balance < 0 else
               '🟢 刚好结清' if balance == 0 else '🟡 预缴 %d 堂' % balance)
        if odd:
            tag += '  ⚠️'
            odd_ones.append((klass, name, paid_count, lessons, hint))
        report.append('%-10s %-12s %5d %5d %6d  %s'
                      % (klass, name, paid_count, charged, balance, tag))
        if balance < 0:
            owed += -balance
    report += ['=' * 62,
               '合计欠 %d 堂' % owed]

    if odd_ones:
        report += ['',
                   '⚠️ 缴费是「涂满 %d 个 1 算一期」,所以已缴堂数应该是 %d 的倍数。'
                   % (args.package, args.package),
                   '   下面这几位不是,代表底色可能多涂或少涂 —— 请核对后用',
                   '   --paid-override 姓名=堂数 重跑一次:']
        for klass, name, paid_count, lessons, hint in odd_ones:
            report.append('     %-10s %-12s 目前算出 %d 堂' % (klass, name, paid_count))
            if hint:
                low, low_date, high, high_date = hint
                report.append('       ↓ 少涂到 %d 堂 → 第 %d 个 1 在 %s'
                              % (low, low, low_date or '(没有这一堂)'))
                report.append('       ↑ 多涂到 %d 堂 → 第 %d 个 1 在 %s'
                              % (high, high, high_date or '(还没上到这一堂)'))

    unknown = set(overrides) - used_overrides
    if unknown:
        report += ['', '⚠️ --paid-override 里这些名字在档案中找不到:%s'
                   % '、'.join(sorted(unknown))]

    report += ['',
               '学生      %d 位' % len(students),
               '课程记录  %d 笔' % len(sessions),
               '付款记录  %d 笔' % len(payments),
               '',
               '⚠️ 上面的「已缴」完全来自旧档的底色。汇入前请先看一眼这张表,',
               '   数字跟你印象中差太多的话,先确认底色的意思有没有被我猜错。',
               '',
               '⚠️ 旧档的 0 当成「请假(免扣)」—— 也就是不扣堂数。',
               '   如果你的 0 其实是「缺席但照算」,重跑一次并加上 --zero-charged。',
               '',
               '下一步:把每个 CSV 汇入对应的工作表(档案 → 汇入 → 在指定储存格取代资料 → A1),',
               '        然后执行 📚 补习管理 → 🔄 刷新总览。',
               '        学生的 WhatsApp 号码和家长称呼旧档没有,要自己补。']

    with open(os.path.join(args.out, '汇入报告.txt'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(report) + '\n')

    print('\n'.join(report))
    print('\n✅ 输出到 %s/' % args.out)


if __name__ == '__main__':
    main()
