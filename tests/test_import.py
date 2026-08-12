#!/usr/bin/env python3
"""验证旧 Excel 汇入时的日期修正逻辑。跑法:python3 tests/test_import.py"""

import datetime
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'tools'))
from import_excel import (parse_header_date, swap_day_month, from_serial,  # noqa: E402
                          read_sheet, fill_rgb, detect_paid_fill, CHARGEABLE, YES)

D = datetime.date
passed = failed = 0


def check(label, actual, expected):
    global passed, failed
    ok = actual == expected
    print(('  ✅ ' if ok else '  ❌ ') + label +
          ('  → %s' % (actual,) if ok else '\n       得到 %s\n       预期 %s' % (actual, expected)))
    if ok:
        passed += 1
    else:
        failed += 1


print('\n── 日月对调 ──')
check('2026-12-01 → 2026-01-12', swap_day_month(D(2026, 12, 1)), D(2026, 1, 12))
check('2026-09-03 → 2026-03-09', swap_day_month(D(2026, 9, 3)), D(2026, 3, 9))
check('日 > 12 无法对调', swap_day_month(D(2026, 1, 19)), None)
check('日月相同 → 不变', swap_day_month(D(2026, 2, 2)), D(2026, 2, 2))

print('\n── Excel 序号解码 ──')
check('46357 = 2026-12-01', from_serial(46357), D(2026, 12, 1))
check('46055 = 2026-02-02', from_serial(46055), D(2026, 2, 2))

print('\n── 标题日期解析(开启修正)──')
# 这几个值直接取自使用者的旧档,预期结果由「同班星期一致」推得
cases = [
    (46357, D(2026, 1, 12), '初二科学 周一'),
    (46268, D(2026, 3, 9), '初二科学 周一'),
    (46055, D(2026, 2, 2), '初二科学 周一(日月相同)'),
    (46144, D(2026, 2, 5), '高二补习 周四'),
    (46175, D(2026, 2, 6), '高三 周五'),
    (46025, D(2026, 3, 1), '高一 周日'),
]
for raw, expect, why in cases:
    got, _ = parse_header_date(raw, swap_fix=True)
    check('%s → %s  (%s)' % (raw, expect, why), got, expect)

print('\n── 文字日期一律照 日/月/年 读 ──')
for text, expect in [('19/1/2026', D(2026, 1, 19)),
                     ('26/1/2026', D(2026, 1, 26)),
                     ('16/1/26', D(2026, 1, 16)),
                     ('2026-03-16', D(2026, 3, 16))]:
    got, _ = parse_header_date(text, swap_fix=True)
    check('「%s」' % text, got, expect)

print('\n── 关闭修正时应保留原值 ──')
got, _ = parse_header_date(46357, swap_fix=False)
check('--no-swap-fix 时 46357 维持 2026-12-01', got, D(2026, 12, 1))

print('\n── 坏资料 ──')
check('空白', parse_header_date(None, True)[0], None)
check('乱字串', parse_header_date('第三週', True)[0], None)


# ── 端到端:组一张假的旧点名表,验证 1/0 与底色的解读 ────────────────
print('\n── 端到端:1=算钱、0=不算钱、底色=已缴 ──')
import openpyxl                                                    # noqa: E402
from openpyxl.styles import PatternFill                            # noqa: E402

YELLOW = PatternFill(fill_type='solid', fgColor='FFFFFF00')
wb = openpyxl.Workbook()
ws = wb.active
ws.title = '测试班'

# 4 个上课日:两个文字日期、两个「被 Excel 吃掉」的数字日期
#   B1 文字 15/1/2026(四)  C1 数字 46144→2026-02-05(四)
#   D1 文字 26/2/2026(四)  E1 数字 46145→2026-03-05(四)
ws['B1'], ws['C1'], ws['D1'], ws['E1'] = '15/1/2026', 46144, '26/2/2026', 46145

# 甲:全出席,前 3 堂涂色         → 已上 4、已缴 3
# 乙:第 2 堂是 0(不算钱),整段涂色 → 已上 3、已缴 3(0 那格不计入已缴)
# 丙:中途才加入(前两格空白)     → 已上 2、已缴 0
ws['A2'], ws['A3'], ws['A4'] = '甲', '乙', '丙'
for col, v in zip('BCDE', [1, 1, 1, 1]):
    ws['%s2' % col] = v
for col, v in zip('BCDE', [1, 0, 1, 1]):
    ws['%s3' % col] = v
for col, v in zip('DE', [1, 1]):
    ws['%s4' % col] = v

for col in 'BCD':
    ws['%s2' % col].fill = YELLOW      # 甲:涂 3 格
for col in 'BCDE':
    ws['%s3' % col].fill = YELLOW      # 乙:涂满 4 格,其中一格是 0

detected, counts = detect_paid_fill(wb)
check('自动侦测底色', detected, 'FFFFFF00')
check('底色格数', counts['FFFFFF00'], 7)
check('没有底色的格子回传 None', fill_rgb(ws['E2']), None)

records, weekdays, bad = read_sheet(ws, ws, True, 'FFFFFF00', False, [])
check('日期全部还原到星期四', dict(weekdays), {'四': 4})
check('没有解析不了的日期', bad, [])


def tally(name):
    used = sum(1 for _, st, _ in records[name] if CHARGEABLE.get(st) == YES)
    paid = sum(1 for _, _, p in records[name] if p)
    return used, paid


check('甲 已上/已缴', tally('甲'), (4, 3))
check('乙 已上/已缴(0 不算钱、也不算已缴)', tally('乙'), (3, 3))
check('丙 已上/已缴(空白不产生记录)', tally('丙'), (2, 0))
check('乙 的 0 被记成请假(免扣)',
      sorted({st for _, st, _ in records['乙']}), ['出席', '请假(免扣)'])

records2, _, _ = read_sheet(ws, ws, True, 'FFFFFF00', True, [])   # --zero-charged
used2 = sum(1 for _, st, _ in records2['乙'] if CHARGEABLE.get(st) == YES)
check('--zero-charged 时乙的 0 改为照算', used2, 4)

print('\n%s%d 项通过\n' % ('❌ %d 项失败,' % failed if failed else '🎉 全数通过,', passed))
sys.exit(1 if failed else 0)
